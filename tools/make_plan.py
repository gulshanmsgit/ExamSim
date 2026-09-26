"""Convert the CST 2026 study planner (Excel) into an ExamSim study-plan pack.

Keeps every row and field of the 'Paper 1 Plan' and 'Paper 2 Plan' sheets (including current-affairs
and revision rows), the subject marks/weights and notes from 'Tracker', and every day of 'Daily Calendar'.
Usage: python tools/make_plan.py packs/gpstr-cst-2026-plan.json
"""
import json, re, sys, warnings
from pathlib import Path
import openpyxl

warnings.filterwarnings("ignore")
SRC = Path(r"C:\Users\DELL\Downloads\CST_2026_Study_Planner.xlsx")
wb = openpyxl.load_workbook(SRC, data_only=True)
ID = re.compile(r"^P[12]-[A-Z0-9]+-\d+$")


def day(v):
    return v.strftime("%Y-%m-%d") if hasattr(v, "strftime") else (str(v)[:10] if v else None)


def text(v):
    return "" if v is None else str(v).strip()


def sheet_plan(name):
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    out = {"title": text(rows[0][0]), "source": text(rows[1][0]), "howTo": text(rows[2][0]), "groups": []}
    group = None
    for r in rows[4:]:
        first = r[0]
        if isinstance(first, str) and ID.match(first.strip()):
            group["rows"].append({
                "id": first.strip(), "subject": text(r[1]), "topic": text(r[2]), "coverage": text(r[3]), "priority": text(r[4]),
                "hours": float(r[5]) if r[5] not in (None, "") else 0, "start": day(r[6]), "target": day(r[7]), "day": text(r[8]),
                "status": text(r[9]) or "Not Started", "rev1": text(r[10]) or "No", "rev2": text(r[11]) or "No",
                "confidence": int(r[12]) if isinstance(r[12], (int, float)) else None, "tip": text(r[13]), "notes": text(r[14])})
        elif isinstance(first, str) and first.strip():
            parts = [p.strip() for p in first.split("|")]
            group = {"name": parts[0], "meta": [p for p in parts[1:] if p], "rows": []}
            out["groups"].append(group)
    out["groups"] = [g for g in out["groups"] if g["rows"]]
    return out


p1, p2 = sheet_plan("Paper 1 Plan"), sheet_plan("Paper 2 Plan")

# Tracker: exam date, marks/weights per subject, and the explanatory notes
tr = list(wb["Tracker"].iter_rows(values_only=True))
exam_date, marks, notes = None, {}, []
for r in tr:
    if exam_date is None and any(hasattr(c, "strftime") for c in r):
        exam_date = day(next(c for c in r if hasattr(c, "strftime")))
    if r[0] in ("Paper 1", "Paper 2") and r[1] and isinstance(r[2], (int, float)):
        marks[text(r[1])] = r[2]
    if isinstance(r[0], str) and re.match(r"^\d+\.\s", r[0]):
        notes.append(text(r[0]))

cal = []
for r in list(wb["Daily Calendar"].iter_rows(values_only=True))[3:]:
    if not hasattr(r[0], "strftime"):
        continue
    split = lambda v: [t.strip() for t in text(v).split("/") if t.strip()]
    cal.append({"date": day(r[0]), "day": text(r[1]), "phase": text(r[2]), "hours": r[3] or 0,
                "p2": split(r[4]), "p1": split(r[5]), "habit": text(r[6]), "done": text(r[7]) == "Yes"})

pack = {
    "examsimPack": 1,
    "name": "Study plan · CST 2026 (from your Excel planner)",
    "description": "Your complete planner: Paper 1 and Paper 2 plan sheets (every topic with priority, hours, dates, status, revisions, confidence, tips), "
                   "the tracker and the day-by-day calendar. Opens in the 📅 Study plan tab.",
    "subjects": [],
    "studyPlan": {"title": "CST 2026 study plan", "examDate": exam_date, "marks": marks, "notes": notes,
                  "statuses": ["Not Started", "In Progress", "Completed", "Review Required"],
                  "sheets": {"p1": p1, "p2": p2}, "calendar": cal},
}
Path(sys.argv[1]).write_text(json.dumps(pack, ensure_ascii=False, indent=1), encoding="utf-8")
n1 = sum(len(g["rows"]) for g in p1["groups"]); n2 = sum(len(g["rows"]) for g in p2["groups"])
print(f"exam {exam_date} | Paper 1: {len(p1['groups'])} groups, {n1} rows | Paper 2: {len(p2['groups'])} groups, {n2} rows | calendar {len(cal)} days | marks {len(marks)} | notes {len(notes)}")
for g in p1["groups"] + p2["groups"]:
    print(f"  {len(g['rows']):3d}  {g['name']}  {g['meta']}")
